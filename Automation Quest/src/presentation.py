from __future__ import annotations

import xml.etree.ElementTree as element_tree
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from src.intake import IntakeSummary


def create_presentation(
    intake: IntakeSummary,
    presentation_path: Path,
    junit_path: Path,
    command: str,
) -> None:
    presentation = Presentation()
    _add_title_slide(presentation, "AI Quality Navigator", "Competition execution briefing")
    _add_bullets(
        presentation,
        "Input provenance",
        [
            f"Source: {intake.source_name}",
            f"SHA-256: {intake.source_sha256}",
            "AI-generated interpretation requires human review before automation.",
        ],
    )
    _add_bullets(
        presentation,
        "Risk-prioritized intent",
        [f"{item.priority} | {item.identifier} | {item.title}" for item in intake.requirements[:12]],
    )
    _add_bullets(presentation, "Execution", [f"Command: {command}", *_junit_summary(junit_path)])
    _add_bullets(presentation, "Quality intelligence", _quality_intelligence(junit_path.parent))
    _add_bullets(
        presentation,
        "Evidence and next actions",
        [
            "Review HTML dashboard, JUnit XML, screenshots, traces, and video before sharing.",
            "Classify failed tests before treating any failure as a potential product defect.",
            "Automate only journeys verified against the supplied application and requirements.",
        ],
    )
    presentation_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(presentation_path)


def _junit_summary(junit_path: Path) -> list[str]:
    if not junit_path.is_file():
        return ["Test results are not available; pytest did not produce JUnit XML."]
    root = element_tree.parse(junit_path).getroot()
    suite = root if root.tag == "testsuite" else root.find("testsuite")
    if suite is None:
        return ["JUnit XML has no test suite summary."]
    tests = suite.attrib.get("tests", "0")
    failures = suite.attrib.get("failures", "0")
    errors = suite.attrib.get("errors", "0")
    skipped = suite.attrib.get("skipped", "0")
    return [f"Tests: {tests}", f"Failures: {failures}", f"Errors: {errors}", f"Skipped: {skipped}"]


def _quality_intelligence(reports_path: Path) -> list[str]:
    path = reports_path / "quality-dashboard.json"
    if not path.is_file():
        return ["Quality dashboard is not available for this execution."]
    dashboard = json.loads(path.read_text(encoding="utf-8"))
    recommendations = dashboard.get("recommendations", [])[:3]
    return [
        f"Risk score: {dashboard.get('risk_score', 0)}/100",
        f"Safe locator-healing events: {dashboard.get('healing', {}).get('safe_events', 0)}",
        *[f"{item['priority']} | {item['action']}" for item in recommendations],
    ]


def _add_title_slide(presentation: Presentation, title: str, subtitle: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def _add_bullets(presentation: Presentation, title: str, entries: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for index, entry in enumerate(entries):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = entry
        paragraph.font.size = Pt(18)
        paragraph.space_after = Pt(8)